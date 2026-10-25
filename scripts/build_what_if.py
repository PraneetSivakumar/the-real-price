"""Phase 6 - builds docs/what_if_simulator.xlsx: change the yellow cells, and the savings,
hours of work and savings-runway date recalculate with live Excel formulas.
Baselines come from the enriched data (post-internship behaviour, 1 Jul - 25 Oct = 117 days).
Run from the project root:  python3 scripts/build_what_if.py
"""
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

df = pd.read_csv("data/processed/transactions_enriched.csv")
m = json.load(open("dashboard/metrics.json"))
sp = df[df.direction == "Debit"]
after = sp[sp.period == "After internship (Jul-Oct)"]
late = (after.hour >= 21) | (after.hour < 5)
lt = after[(after.subcategory == "Takeaway delivery") & (after.shift_day) & late]
base = {
    "late_take_spend": round(lt.amount_gbp.sum(), 2),
    "late_take_orders": int(len(lt)),
    "topup_spend": round(after[after.subcategory == "Top-up shop"].amount_gbp.sum(), 2),
    "late_online": round(after[after.channel.isin(["Online", "App"]) & ((after.hour >= 22) | (after.hour < 5))
                               & after.category.isin(["Clothing", "Shopping (misc)", "Entertainment"])].amount_gbp.sum(), 2),
    "tub_spend": round(sp[sp.subcategory == "Protein (small tub, in-store)"].amount_gbp.sum(), 2),
    "zombie_monthly": 5.99 + 1.99,
}

YEL = PatternFill("solid", fgColor="FFF2B3")
HEAD = Font(bold=True, size=13)
wb = Workbook()
ws = wb.active
ws.title = "Simulator"
ws.column_dimensions["A"].width = 62
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 60

def row(r, label, value, note="", fmt=None, inp=False, bold=False):
    ws.cell(r, 1, label).font = Font(bold=bold)
    c = ws.cell(r, 2, value)
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = YEL
    if bold:
        c.font = Font(bold=True)
    ws.cell(r, 3, note).alignment = Alignment(wrap_text=True)

ws["A1"] = "The Real Price — what-if savings simulator (SIMULATED data)"; ws["A1"].font = Font(bold=True, size=15)
ws["A2"] = "Change the yellow cells. Everything else recalculates."
ws["A4"] = "1. Your settings"; ws["A4"].font = HEAD
row(5, "Hourly wage (£)", 12.71, "UK National Living Wage from April 2026", "£0.00", True)
row(6, "Days from 25 Oct to 31 Dec", 67, "", "0", True)
row(7, "Cancel Disney+ and Google One? (Y/N)", "Y", "Zombie / duplicate subscriptions", None, True)
row(8, "Share of late shift-day takeaways replaced by meal prep", 0.75, "", "0%", True)
row(9, "Cost of a prepared meal (£)", 3.00, "", "£0.00", True)
row(10, "Share of top-up trips moved into the weekly shop", 0.50, "", "0%", True)
row(11, "How much cheaper the weekly shop is than a top-up", 0.25, "Assumption — UK Basket Index will measure it", "0%", True)
row(12, "Share of late-night online shopping avoided (24-hour rule)", 0.50, "", "0%", True)
row(13, "Buy protein in bulk only? (Y/N)", "Y", "", None, True)
dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False); ws.add_data_validation(dv); dv.add("B7"); dv.add("B13")

ws["A15"] = "2. Baseline from the data (after the internship: 117 days)"; ws["A15"].font = HEAD
row(16, "Late takeaways on shift days — spend (£)", base["late_take_spend"], "1 Jul – 25 Oct", "£#,##0.00")
row(17, "Late takeaways on shift days — orders", base["late_take_orders"], "", "0")
row(18, "Top-up shops — spend (£)", base["topup_spend"], "", "£#,##0.00")
row(19, "Late-night online clothing / shopping / games (£)", base["late_online"], "22:00–05:00", "£#,##0.00")
row(20, "In-store protein tubs, whole period (£, 298 days)", base["tub_spend"], "", "£#,##0.00")
row(21, "Zombie subscriptions per month (£)", base["zombie_monthly"], "Disney+ £5.99 + Google One £1.99", "£0.00")
row(22, "Bulk protein cost vs in-store tub (per gram)", round(m["protein_per_20g"]["Protein (bulk online)"] / m["protein_per_20g"]["Protein (small tub, in-store)"], 3), "From the protein economics analysis", "0%")
row(23, "Current balance on 25 Oct (£)", m["end_balance"], "", "£#,##0.00")
row(24, "Net cash flow per day after the internship (£)", m["runway"]["net_per_day"], "Negative = savings going down", "£0.00")

ws["A26"] = "3. Savings per day (£)"; ws["A26"].font = HEAD
row(27, "Cancel zombie subscriptions", '=IF(B7="Y",B21*12/365,0)', "", "£0.00")
row(28, "Shift-day meal prep", "=B8*MAX(B16-B9*B17,0)/117", "", "£0.00")
row(29, "Move top-ups into the weekly shop", "=B10*B18*B11/117", "", "£0.00")
row(30, "Protein: bulk only", '=IF(B13="Y",B20*(1-B22)/298,0)', "", "£0.00")
row(31, "24-hour rule for late-night online shopping", "=B12*B19/117", "", "£0.00")
row(32, "Total saved per day", "=SUM(B27:B31)", "", "£0.00", bold=True)

ws["A34"] = "4. Results"; ws["A34"].font = HEAD
row(35, "Saved by 31 Dec 2026 (£)", "=B32*B6", "", "£#,##0", bold=True)
row(36, "Saved per year (£)", "=B32*365", "", "£#,##0", bold=True)
row(37, "Hours of work saved per year", "=B36/B5", "", "0", bold=True)
row(38, "Savings run out (no change)", '=IF(B24<0,DATE(2026,10,25)+B23/-B24,"Not running out")', "Date the balance hits £0 at today's rate", "d mmm yyyy")
row(39, "Savings run out (with the plan)", '=IF(B24+B32<0,DATE(2026,10,25)+B23/-(B24+B32),"Not running out")', "", "d mmm yyyy", bold=True)
ws["A41"] = "Note: the plan slows the fall; only new income (the job search) stops it."
ws["A41"].font = Font(italic=True)
wb.save("docs/what_if_simulator.xlsx")
print("docs/what_if_simulator.xlsx written", base)
